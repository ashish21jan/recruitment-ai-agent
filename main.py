"""
Recruitment AI Agent - FastAPI Application
This application uses Google Gemini API to evaluate resumes against job descriptions
and generate personalized emails for candidates.
"""

import os
import io
from typing import List, Optional
from fastapi import FastAPI, File, UploadFile, Form, HTTPException, Request
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
import google.generativeai as genai
from pypdf import PdfReader
from docx import Document
from dotenv import load_dotenv
import json

# Load environment variables
load_dotenv()

# Configure Gemini API
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
if not GEMINI_API_KEY:
    raise ValueError("GEMINI_API_KEY environment variable is not set")

genai.configure(api_key=GEMINI_API_KEY)
model = genai.GenerativeModel('gemini-2.5-flash')

# Initialize FastAPI app
app = FastAPI(
    title="Recruitment AI Agent",
    description="AI-powered recruitment assistant for resume evaluation and candidate matching",
    version="1.0.0"
)

# Mount static files and templates
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")


# Pydantic models for API requests
class GenerateJDRequest(BaseModel):
    """Request model for AI-generated job description"""
    job_title: str
    years_of_experience: int
    must_have_skills: str
    company_name: str
    employment_type: str
    industry: str
    location: str


class CandidateEvaluation(BaseModel):
    """Model for candidate evaluation results"""
    filename: str
    score: int
    missing_skills: List[str]
    remarks: str


# Utility functions for text extraction
def extract_text_from_pdf(file_content: bytes) -> str:
    """
    Extract text content from a PDF file.
    
    Args:
        file_content: Bytes content of the PDF file
        
    Returns:
        Extracted text as string
    """
    try:
        pdf_reader = PdfReader(io.BytesIO(file_content))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error extracting PDF: {str(e)}")


def extract_text_from_docx(file_content: bytes) -> str:
    """
    Extract text content from a DOCX file.
    
    Args:
        file_content: Bytes content of the DOCX file
        
    Returns:
        Extracted text as string
    """
    try:
        doc = Document(io.BytesIO(file_content))
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text.strip()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error extracting DOCX: {str(e)}")


def extract_text_from_file(file_content: bytes, filename: str) -> str:
    """
    Extract text from a file based on its extension.
    
    Args:
        file_content: Bytes content of the file
        filename: Name of the file with extension
        
    Returns:
        Extracted text as string
    """
    if filename.lower().endswith('.pdf'):
        return extract_text_from_pdf(file_content)
    elif filename.lower().endswith(('.docx', '.doc')):
        return extract_text_from_docx(file_content)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {filename}")


# API Endpoints
@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    """
    Render the home page with the job description input form.
    """
    return templates.TemplateResponse("index.html", {"request": request})


@app.post("/generate-jd")
async def generate_job_description(jd_request: GenerateJDRequest):
    """
    Generate a complete job description using Gemini AI.
    
    Args:
        jd_request: Job description parameters including title, experience, skills, etc.
        
    Returns:
        JSON object containing the generated job description text
    """
    try:
        prompt = f"""
        Generate a professional and comprehensive job description with the following details:
        
        Job Title: {jd_request.job_title}
        Years of Experience: {jd_request.years_of_experience}
        Must-Have Skills: {jd_request.must_have_skills}
        Company Name: {jd_request.company_name}
        Employment Type: {jd_request.employment_type}
        Industry: {jd_request.industry}
        Location: {jd_request.location}
        
        Please create a detailed job description that includes:
        - A brief company overview
        - Job responsibilities (5-7 bullet points)
        - Required qualifications and skills
        - Preferred qualifications
        - Benefits (if applicable)
        
        Make it professional and engaging.
        """
        
        response = model.generate_content(prompt)
        jd_text = response.text
        
        return {"job_description": jd_text}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generating job description: {str(e)}")


@app.post("/extract-jd")
async def extract_job_description(jd_file: UploadFile = File(...)):
    """
    Extract job description text from an uploaded PDF or DOCX file.
    
    Args:
        jd_file: Uploaded job description file
        
    Returns:
        JSON object containing the extracted job description text
    """
    try:
        file_content = await jd_file.read()
        jd_text = extract_text_from_file(file_content, jd_file.filename)
        
        return {"job_description": jd_text}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error extracting job description: {str(e)}")


@app.post("/evaluate", response_class=HTMLResponse)
async def evaluate_candidates(
    request: Request,
    job_description: str = Form(...),
    resumes: List[UploadFile] = File(...)
):
    """
    Evaluate multiple candidate resumes against a job description.
    
    Args:
        job_description: The job description text
        resumes: List of uploaded resume files (up to 10)
        
    Returns:
        HTML page with evaluation results, including scores, missing skills,
        and auto-generated emails for the best and rejected candidates
    """
    try:
        # Validate number of resumes
        if len(resumes) > 10:
            raise HTTPException(status_code=400, detail="Maximum 10 resumes allowed")
        
        if len(resumes) == 0:
            raise HTTPException(status_code=400, detail="At least one resume is required")
        
        # Evaluate each resume
        evaluations = []
        
        for resume_file in resumes:
            try:
                # Extract resume text
                resume_content = await resume_file.read()
                resume_text = extract_text_from_file(resume_content, resume_file.filename)
                
                # Create prompt for Gemini to evaluate the resume
                evaluation_prompt = f"""
                You are a professional recruiter. Evaluate the following resume against the job description.
                
                Job Description:
                {job_description}
                
                Resume:
                {resume_text}
                
                Provide your evaluation in the following JSON format ONLY (no additional text):
                {{
                    "score": <integer from 0 to 100>,
                    "missing_skills": [<list of key skills from JD not found in resume>],
                    "remarks": "<one-sentence explanation for the score>"
                }}
                
                Be strict in your evaluation. Consider experience level, skills match, and overall fit.
                """
                
                # Get evaluation from Gemini
                response = model.generate_content(evaluation_prompt)
                response_text = response.text.strip()
                
                # Extract JSON from response (handle code blocks if present)
                if "```json" in response_text:
                    response_text = response_text.split("```json")[1].split("```")[0].strip()
                elif "```" in response_text:
                    response_text = response_text.split("```")[1].split("```")[0].strip()
                
                evaluation_data = json.loads(response_text)
                
                # Create evaluation object
                evaluation = CandidateEvaluation(
                    filename=resume_file.filename,
                    score=evaluation_data["score"],
                    missing_skills=evaluation_data["missing_skills"],
                    remarks=evaluation_data["remarks"]
                )
                evaluations.append(evaluation)
                
            except Exception as e:
                print(f"Error evaluating {resume_file.filename}: {str(e)}")
                # Add a default evaluation for files that fail
                evaluations.append(CandidateEvaluation(
                    filename=resume_file.filename,
                    score=0,
                    missing_skills=["Error processing resume"],
                    remarks=f"Could not process this resume: {str(e)}"
                ))
        
        # Find the best candidate
        best_candidate = max(evaluations, key=lambda x: x.score)
        
        # Generate personalized interview invitation email for best candidate
        interview_email_prompt = f"""
        Generate a professional and warm interview invitation email for the following candidate.
        
        Job Description:
        {job_description}
        
        Candidate's Resume Score: {best_candidate.score}/100
        Candidate's Filename: {best_candidate.filename}
        
        Create a personalized email that:
        - Congratulates them on being selected for an interview
        - Mentions their strong qualifications
        - Invites them to schedule an interview
        - Is warm and professional in tone
        
        Format the email with proper greeting and signature placeholders.
        """
        
        interview_email_response = model.generate_content(interview_email_prompt)
        interview_email = interview_email_response.text
        
        # Generate rejection email template
        rejection_email_prompt = f"""
        Generate a professional and empathetic rejection email for candidates who were not selected.
        
        Job Description:
        {job_description}
        
        Create a polite rejection email that:
        - Thanks them for their interest and time
        - Informs them they were not selected for this position
        - Encourages them to apply for future opportunities
        - Is respectful and professional in tone
        
        Format the email with proper greeting and signature placeholders.
        """
        
        rejection_email_response = model.generate_content(rejection_email_prompt)
        rejection_email = rejection_email_response.text
        
        # Sort evaluations by score (descending)
        evaluations_sorted = sorted(evaluations, key=lambda x: x.score, reverse=True)
        
        # Render results page
        return templates.TemplateResponse("results.html", {
            "request": request,
            "evaluations": evaluations_sorted,
            "best_candidate": best_candidate,
            "interview_email": interview_email,
            "rejection_email": rejection_email,
            "job_description": job_description
        })
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error evaluating candidates: {str(e)}")


@app.get("/health")
async def health_check():
    """
    Health check endpoint to verify the API is running.
    """
    return {"status": "healthy", "message": "Recruitment AI Agent is running"}


if __name__ == "__main__":
    import uvicorn
    port = int(os.getenv("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)


